from __future__ import annotations
import os, sys, re, json, math, csv, statistics, hashlib, copy, shutil, importlib.util
from pathlib import Path
from xml.sax.saxutils import escape

ROOT = Path(__file__).resolve().parent
PROJECT = next(p for p in ROOT.parents if (p/'unified_actor3d').is_dir())
TMP = ROOT / '_work' / 'tmp'
TMP.mkdir(parents=True, exist_ok=True)
os.environ['TEMP'] = os.environ['TMP'] = str(TMP)
os.environ['PYTHONIOENCODING'] = 'utf-8'
sys.dont_write_bytecode = True

from PIL import Image
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.colors import HexColor, Color, white, black
from pdf2image import convert_from_path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

RUNTIME = Path(r'C:\Program Files\WindowsApps\OpenAI.CodexPrimaryRuntime.v26-1025-760-0_26.1025.760.0_x64__3k8sg7r9htsxt\dependencies')
POPPLER = RUNTIME / 'native' / 'poppler' / 'Library' / 'bin'
FIG = ROOT / 'figures'
FIG.mkdir(exist_ok=True)
pdfmetrics.registerFont(TTFont('CN', r'C:\Windows\Fonts\simhei.ttf'))
BLUE, TEAL, ORANGE, RED, GRAY = map(HexColor, ['#285477','#247C73','#C36A27','#B33A3A','#64717B'])
LIGHT = HexColor('#E8EDF0')

EVIDENCE = {
 'E01': ('2.3.4冻结源码和参数', PROJECT/'交付文件/ADP排查与修复_20261012/current_source'),
 'E02': ('ADP原因排查 修复与可解释性验证', PROJECT/'交付文件/ADP排查与修复_20261012/报告与图表/README.md'),
 'E03': ('安全层同图对照分析', PROJECT/'交付文件/安全层同图对照分析_20261028/分析结果/README.md'),
 'E04': ('APF切向退化动态检查', PROJECT/'交付文件/APF切向退化动态检查_20261029/现象与依据.md'),
 'E05': ('噪声与外扰鲁棒性实验', PROJECT/'交付文件/噪声与外扰鲁棒性实验_20261029/README.md'),
 'E06': ('神经ADP训练模块与单机自动脱困', PROJECT/'results/neural_adp_supervisor/README_自动脱困验证.md'),
 'E07': ('五机规划与脱困策略六组合对照', PROJECT/'results/five_uav_comparison3d/campaign_20261017_210011/README_对照实验.md'),
 'E08': ('冻结Actor五机接入与队形诊断', PROJECT/'results/five_uav_actor3d/README_五机验证.md'),
 'E09': ('未知地图R34性能与ADP交接', PROJECT/'docs/HANDOFF_PERFORMANCE_ADP_20261027.md'),
 'E10': ('已知地图历史实现与性能归档', PROJECT/'README_LEGACY_FULL_MAP.md'),
 'E11': ('冻结Actor单机CPU时钟与故障检查', PROJECT/'results/cpu_realtime_adp3d/README_CPU实时性验证.md'),
 'E12': ('未知地图体系与接口', PROJECT/'docs/UNKNOWN_MAP_ARCHITECTURE.md'),
 'E13': ('TD3神经网络结构与训练说明', PROJECT/'neural_adp3d/README_训练与运行.md'),
}

def read_csv(path):
    with open(path, encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))

def fval(row, key): return float(row[key])

ROBUST = PROJECT / '交付文件/噪声与外扰鲁棒性实验_20261029'
DYNAMIC = PROJECT / '交付文件/APF切向退化动态检查_20261029'
SAFETY = PROJECT / '交付文件/安全层同图对照分析_20261028/分析结果'

def figure_file(name, height, draw):
    pdf = FIG / (name+'.pdf')
    c = canvas.Canvas(str(pdf), pagesize=(504,height))
    c.setTitle(name)
    c.setFillColor(white);c.rect(0,0,504,height,fill=1,stroke=0)
    c.setFont('CN',10);draw(c,504,height);c.showPage();c.save()
    im = convert_from_path(str(pdf), dpi=220, poppler_path=str(POPPLER), first_page=1,last_page=1)[0]
    out=FIG/(name+'.png');im.save(out)
    return out

def txt(c,x,y,s,size=10,color=black,align='left'):
    s=str(s).replace('²','^2').replace('³','^3').replace('−','-')
    c.setFont('CN',size);c.setFillColor(color)
    getattr(c, {'left':'drawString','center':'drawCentredString','right':'drawRightString'}[align])(x,y,str(s))

def line(c, x1,y1,x2,y2,color=GRAY,width=1,dash=None):
    c.setStrokeColor(color);c.setLineWidth(width);c.setDash(dash or [])
    c.line(x1,y1,x2,y2);c.setDash([])

def arrow(c,x1,y1,x2,y2,color=BLUE):
    line(c,x1,y1,x2,y2,color,1.2)
    a=math.atan2(y2-y1,x2-x1)
    p=c.beginPath();p.moveTo(x2,y2)
    p.lineTo(x2-7*math.cos(a-.4),y2-7*math.sin(a-.4))
    p.lineTo(x2-7*math.cos(a+.4),y2-7*math.sin(a+.4));p.close()
    c.setFillColor(color);c.drawPath(p,fill=1,stroke=0)

def box(c,x,y,w,h,labels):
    c.setFillColor(HexColor('#F5F7F8'));c.setStrokeColor(HexColor('#9FAEB7'))
    c.rect(x,y,w,h,stroke=1,fill=1)
    for i,s in enumerate(labels): txt(c,x+w/2,y+h/2+(len(labels)-1)*7-i*14,s,10,BLUE,'center')

def axes(c,rect,xr,yr,xt,yt,xlabel,ylabel):
    x,y,w,h=rect
    X=lambda v:x+(v-xr[0])/(xr[1]-xr[0])*w
    Y=lambda v:y+(v-yr[0])/(yr[1]-yr[0])*h
    for v in yt:
        line(c,x,Y(v),x+w,Y(v),LIGHT,.5)
        label=(f'{v:.2g}' if 0<abs(v)<.15 else f'{v:.1f}'.rstrip('0').rstrip('.'))
        txt(c,x-6,Y(v)-3,label,8,GRAY,'right')
    for v in xt:
        line(c,X(v),y,X(v),y-3,GRAY,.5)
        label=(f'{v:.2g}' if 0<abs(v)<.15 else f'{v:.1f}'.rstrip('0').rstrip('.'))
        txt(c,X(v),y-15,label,8,GRAY,'center')
    line(c,x,y,x+w,y,GRAY,.7);line(c,x,y,x,y+h,GRAY,.7)
    txt(c,x+w/2,y-31,xlabel,9,GRAY,'center')
    txt(c,x,y+h+9,ylabel,9,GRAY)
    return X,Y

def polyline(c, points, color, width=1.25):
    if len(points)<2:return
    path=c.beginPath();path.moveTo(*points[0])
    for p in points[1:]:path.lineTo(*p)
    c.setStrokeColor(color);c.setLineWidth(width);c.setDash([]);c.drawPath(path)

def make_figures():
    robust=read_csv(ROBUST/'逐任务汇总.csv')
    dynamic=read_csv(DYNAMIC/'逐任务汇总.csv')
    safety=read_csv(SAFETY/'任务结局汇总.csv')
    assert len(robust)==912 and all(r['success']=='True' for r in robust)
    assert len(dynamic)==100 and sum(int(r['success']) for r in dynamic)==88
    qual=[]
    for sigma in [0,.02,.05,.1]:
        rows=[r for r in robust if r['group']=='S' and float(r['sigma'])==sigma and r['kind']==('baseline' if sigma==0 else 'noise')]
        qual.append(statistics.median(float(r['post5_command_slew_rms_mps3']) for r in rows))
    assert abs(qual[0]-9.6789352098)<1e-6
    assert abs(qual[3]-51.2175392353)<1e-6
    facts={'robust_trials':len(robust),'robust_success':sum(r['success']=='True' for r in robust),
       'min_obstacle_m':min(float(r['min_obstacle_m']) for r in robust),
       'min_pair_m':min(float(r['min_pair_m']) for r in robust),
       'max_model_accel':max(float(r['max_model_accel_mps2']) for r in robust),
       'noise_slew_medians':qual,'noise_slew_ratio':qual[-1]/qual[0],
       'dynamic_trials':len(dynamic),'dynamic_success':sum(int(r['success']) for r in dynamic),
       'safety_R':sum(int(r['R_success']) for r in safety),'safety_S':sum(int(r['S_success']) for r in safety)}
    assert facts['safety_R']==145 and facts['safety_S']==149
    (ROOT/'复算指标.json').write_text(json.dumps(facts,ensure_ascii=False,indent=2),encoding='utf-8')

    def flow(c,w,h):
        entries=[(22,142,['任务和全图','起点 目标 约束']),(191,142,['有限层规划','生成长机路线']),(360,142,['编队参考','跟随轨迹和队形']),
                 (360,43,['APF和ADP','生成五机候选指令']),(191,43,['限幅和安全层','预测 修正 核验']),(22,43,['动力学和记录','更新真实状态'])]
        for x,y,labs in entries:box(c,x,y,122,49,labs)
        arrow(c,144,166,191,166);arrow(c,313,166,360,166)
        arrow(c,421,142,421,92);arrow(c,360,67,313,67);arrow(c,191,67,144,67)
        line(c,83,43,83,20,TEAL);line(c,83,20,421,20,TEAL);arrow(c,421,20,421,43,TEAL)
        txt(c,252,3,'反馈给参考生成 控制器和学习模块',9,TEAL,'center')
    figure_file('system_flow',211,flow)

    def layers(c,w,h):
        for i,z in enumerate([5,10,15,20,25]):
            y=36+i*30
            line(c,52,y,453,y,BLUE,.7)
            txt(c,42,y-3,f'{z} m',9,GRAY,'right')
            for x in range(72,450,24):
                c.setFillColor(HexColor('#B8C9D4'));c.circle(x,y,1.4,fill=1,stroke=0)
        c.setFillColor(HexColor('#D9DDDF'));c.setStrokeColor(GRAY)
        c.rect(218,36,72,75,fill=1,stroke=1);c.ellipse(218,101,290,121,fill=1,stroke=1)
        txt(c,254,75,'有限圆柱',9,GRAY,'center')
        path=[(72,66),(168,66),(168,126),(336,126),(336,66),(432,66)]
        polyline(c,path,TEAL,2.4)
        for p in path:c.setFillColor(TEAL);c.circle(*p,2.6,fill=1,stroke=0)
        txt(c,252,189,'五个高度上的节点和安全连接 示意',11,BLUE,'center')
        txt(c,72,14,'起点',10,TEAL,'center');txt(c,432,14,'终点',10,TEAL,'center')
        txt(c,342,171,'跨层边沿竖直方向连接',9,GRAY,'center')
    figure_file('layers',207,layers)

    def formation(c,w,h):
        x0,y0,scale=344,124,19
        line(c,92,y0,430,y0,GRAY,.6,[3,3]);arrow(c,380,y0,440,y0,BLUE)
        txt(c,446,y0-3,'前',10,BLUE)
        coords=[(0,0),(-3,-2.5),(-3,2.5),(-6,-5),(-6,5)]
        colors=[BLUE,TEAL,ORANGE,HexColor('#7465A5'),HexColor('#9B5C6A')]
        for i,((a,b),col) in enumerate(zip(coords,colors),1):
            x,y=x0+a*scale,y0+b*scale
            c.setFillColor(col);c.circle(x,y,5,stroke=0,fill=1)
            txt(c,x+10,y+(12 if i==1 else -3),f'{i}号  ({a:g}, {b:g})',10,col)
            if i>1:line(c,x0,y0,x,y,LIGHT,.8)
        txt(c,70,230,'左',10,GRAY);txt(c,70,15,'右',10,GRAY)
        txt(c,252,252,'以长机为原点的水平坐标 单位 米',10,BLUE,'center')
    figure_file('formation',268,formation)

    def tangent(c,w,h):
        for x0,label,offset in [(135,'严格正对',0),(385,'略有偏置',25)]:
            c.setFillColor(HexColor('#EBEEF0'));c.setStrokeColor(GRAY);c.circle(x0,108,47,fill=1,stroke=1)
            sx,sy=x0-95,108+offset
            c.setFillColor(BLUE);c.circle(sx,sy,3.6,fill=1,stroke=0)
            arrow(c,sx,sy,x0+71,108,ORANGE)
            nx,ny=sx-x0,sy-108;n=math.hypot(nx,ny)
            arrow(c,sx,sy,sx+nx/n*28,sy+ny/n*28,TEAL)
            if offset:arrow(c,sx,sy,sx+14,sy+43,RED)
            txt(c,x0,188,label,11,BLUE,'center')
            txt(c,x0,29,'投影为零' if not offset else '得到沿表面的分量',10,RED,'center')
        txt(c,252,8,'橙色 目标方向    绿色 外法向    红色 切向示意',9,GRAY,'center')
    figure_file('tangent',211,tangent)

    def safety_chart(c,w,h):
        X,Y=axes(c,(47,43,427,144),(-.5,3.5),(0,105),[],[0,25,50,75,100],'','完成比例 %')
        names=['随机100图','困难12图','检查2图','扩展40图']
        nums=[(96,96,100),(8,11,12),(2,2,2),(39,40,40)]
        for j,((r,s,n),name) in enumerate(zip(nums,names)):
            for off,v,col in [(-.15,r,ORANGE),(.15,s,TEAL)]:
                xx=X(j+off);yy=Y(100*v/n)
                c.setFillColor(col);c.rect(xx-10,Y(0),20,yy-Y(0),fill=1,stroke=0)
                txt(c,xx,yy+5,f'{v}/{n}',8,col,'center')
            txt(c,X(j),25,name,9,GRAY,'center')
        txt(c,320,212,'R 无安全层',10,ORANGE);txt(c,414,212,'S 有安全层',10,TEAL)
    figure_file('safety_outcomes',230,safety_chart)

    pressure={}
    for group in ['R','S']:
        rows=read_csv(DYNAMIC/'traces'/f'C001_{group}.csv')
        pressure[group]=[(float(r['t_after_s']),float(r['fleet_swept_pair_distance'])) for r in rows if r['vehicle']=='1' and float(r['t_after_s'])<=8]
    def pressure_chart(c,w,h):
        X,Y=axes(c,(47,45,423,142),(0,8),(.8,4.2),[0,2,4,6,8],[1,2,3,4],'时间 秒','整步最小机间距 米')
        line(c,X(0),Y(1),X(8),Y(1),RED,.9,[4,3])
        txt(c,X(6),Y(1)+7,'硬门槛 1 m',8,RED)
        for g,col in [('S',TEAL),('R',ORANGE)]:polyline(c,[(X(t),Y(d)) for t,d in pressure[g]],col,1.4)
        t,d=pressure['R'][-1];c.setFillColor(ORANGE);c.circle(X(t),Y(d),3,fill=1,stroke=0)
        txt(c,X(t)+10,Y(d)+13,'R 3.70秒结束',9,ORANGE)
        txt(c,310,211,'R 无安全层',10,ORANGE);txt(c,413,211,'S 有安全层',10,TEAL)
        txt(c,47,211,'显示起始8秒',9,GRAY)
    figure_file('pressure_distance',231,pressure_chart)

    exrows=read_csv(ROBUST/'traces/M01_J020_D030_PY_n01_S.csv')
    ex=[r for r in exrows if r['vehicle']=='1' and float(r['t_before_s'])<=11.2]
    def disturbance_chart(c,w,h):
        X,Y=axes(c,(44,43,197,133),(0,11.2),(-.03,.36),[0,4.2,6.2,11.2],[0,.15,.3],'时间 秒','+Y加速度 米每二次方秒')
        pts=[(X(float(r['t_before_s'])),Y(float(r['disturbance_y']))) for r in ex]
        polyline(c,pts,ORANGE,1.5)
        txt(c,143,198,'阵风',11,BLUE,'center')
        X,Y=axes(c,(291,43,196,133),(0,11.2),(-.075,.075),[0,4.2,6.2,11.2],[-.05,0,.05],'时间 秒','1号机Y位置噪声 米')
        polyline(c,[(X(float(r['t_before_s'])),Y(float(r['noise_y']))) for r in ex],TEAL,.8)
        txt(c,389,198,'测量噪声',11,BLUE,'center')
    figure_file('disturbance',220,disturbance_chart)

    def quality(c,w,h):
        X,Y=axes(c,(52,48,409,140),(-.4,3.4),(0,60),[],[0,15,30,45,60],'每轴位置噪声标准差 米','指令变化率RMS 米每三次方秒')
        for j,(label,v) in enumerate(zip(['0','0.02','0.05','0.10'],qual)):
            c.setFillColor(TEAL);c.rect(X(j)-20,Y(0),40,Y(v)-Y(0),fill=1,stroke=0)
            txt(c,X(j),Y(v)+6,f'{v:.2f}',10,TEAL,'center')
            txt(c,X(j),31,label,10,GRAY,'center')
        txt(c,254,212,'同地图相同绝对五秒窗口 逐任务RMS的中位数',10,BLUE,'center')
    figure_file('noise_quality',231,quality)

    def strategies(c,w,h):
        X,Y=axes(c,(48,43,421,137),(-.5,1.5),(0,150),[],[0,50,100,150],'','完成数 每组144图')
        for j,vs in enumerate([[135,135,135],[71,124,122]]):
            for off,v,col,lab in zip([-.22,0,.22],vs,[GRAY,TEAL,ORANGE],['无脱困','规则','Actor']):
                xx=X(j+off);c.setFillColor(col);c.rect(xx-12,Y(0),24,Y(v)-Y(0),fill=1,stroke=0)
                txt(c,xx,Y(v)+5,str(v),9,col,'center')
            txt(c,X(j),26,['完整全局规划','关闭长机全局规划'][j],10,GRAY,'center')
        for x,s,col in [(191,'无额外脱困',GRAY),(306,'规则航点',TEAL),(404,'冻结Actor',ORANGE)]:txt(c,x,208,s,10,col)
    figure_file('strategy_comparison',230,strategies)

    def unknown(c,w,h):
        entries=[(15,142,['左右目图像','有限视野观察']),(192,142,['匹配和深度','时戳 有效掩码']),(369,142,['每机在线地图','自由 占据 未知']),
                 (369,44,['自由区和前沿','规划与观察机位']),(192,44,['APF和编队','候选加速度']),(15,44,['联合运动守护','执行和制动'])]
        for x,y,l in entries:box(c,x,y,120,50,l)
        arrow(c,135,167,192,167);arrow(c,312,167,369,167);arrow(c,429,142,429,94)
        arrow(c,369,69,312,69);arrow(c,192,69,135,69)
        line(c,75,44,75,22,TEAL);line(c,75,22,253,22,TEAL);arrow(c,253,22,253,44,TEAL)
        txt(c,252,3,'真实几何只进入成像与独立检查',9,GRAY,'center')
    figure_file('unknown_flow',211,unknown)

    traj=read_csv(ROBUST/'traces/M01_B_S.csv')
    colors=[BLUE,TEAL,ORANGE,HexColor('#7465A5'),HexColor('#9B5C6A')]
    def trajectory(c,w,h):
        for xx,axis,yr,yt,title in [(39,'y',(20,60),[20,40,60],'俯视 XY'),(294,'z',(-5,35),[0,15,30],'侧视 XZ')]:
            X,Y=axes(c,(xx,46,198,136.5517),(12,70),yr,[20,40,60],yt,'X 米',axis.upper()+' 米')
            c.setFillColor(HexColor('#E4E7E9'));c.setStrokeColor(HexColor('#BAC2C7'))
            center=40 if axis=='y' else 15
            c.ellipse(X(36),Y(center-4),X(44),Y(center+4),fill=1,stroke=1)
            for i,col in enumerate(colors,1):
                rr=[r for r in traj if int(r['vehicle'])==i]
                pts=[(X(float(r['x'])),Y(float(r[axis]))) for r in rr]
                polyline(c,pts,col,1.0 if i>1 else 1.6)
                c.setFillColor(col);c.circle(*pts[0],2.0,fill=1,stroke=0)
                c.rect(pts[-1][0]-1.8,pts[-1][1]-1.8,3.6,3.6,fill=1,stroke=0)
            txt(c,xx+99,209,title,11,BLUE,'center')
        for i,col in enumerate(colors,1):txt(c,48+90*(i-1),233,f'{i}号机',9,col)
        txt(c,254,6,'圆点 起始   方点 终止   灰色 球体投影',9,GRAY,'center')
    figure_file('actual_trajectory',251,trajectory)
    return facts

# Native Word equations generated through the Office MathML to OMML transform.
def mi(x):return '<mi>'+escape(str(x))+'</mi>'
def mn(x):return '<mn>'+str(x)+'</mn>'
def mo(x):return '<mo>'+escape(str(x))+'</mo>'
def mt(x):return '<mtext>'+escape(str(x))+'</mtext>'
def row(*x):return '<mrow>'+''.join(x)+'</mrow>'
def sub(a,b):return '<msub>'+a+b+'</msub>'
def sup(a,b):return '<msup>'+a+b+'</msup>'
def frac(a,b):return '<mfrac>'+a+b+'</mfrac>'
def sqrt(a):return '<msqrt>'+a+'</msqrt>'
def par(a):return row(mo('('),a,mo(')'))
def norm(a):return row(mo('‖'),a,mo('‖'))
def abs_m(a):return row(mo('|'),a,mo('|'))
def idx(a,b):return sub(mi(a),mi(b))
def arr(*rows):return '<mtable columnalign="left">'+''.join('<mtr><mtd>'+r+'</mtd></mtr>' for r in rows)+'</mtable>'
def vec(*items):return '<mfenced open="[" close="]"><mtable>'+''.join('<mtr><mtd>'+r+'</mtd></mtr>' for r in items)+'</mtable></mfenced>'
def seq(*items):return row(*sum(([x,mo(',')] for x in items),[])[:-1])
def sat(L,a):return row(sub(mi('sat'),mi(L)),par(a))
def summand(n,term,var='j',start=1):return row('<munderover>'+mo('∑')+row(mi(var),mo('='),mn(start))+n+'</munderover>',term)

def equations():
    p,v,u,d=map(mi,['p','v','u','d']);dt=mi('Δt')
    pk,vk,uk,dk=map(lambda a:idx(a,'k'),['p','v','u','d'])
    k1=row(mi('k'),mo('+'),mn(1))
    mv=row(sub(v,k1),mo('='),sat('v',row(vk,mo('+'),dt,par(row(uk,mo('+'),dk)))))
    mp=row(sub(p,k1),mo('='),pk,mo('+'),dt,sub(v,k1))
    eq={
      'saturation':row(sub(mi('sat'),mi('L')),par(mi('z')),mo('='),frac(mi('z'),row(mi('max'),par(seq(mn(1),frac(norm(mi('z')),mi('L'))))))),
      'motion_student':arr(mv,mp),
      'motion':arr(row(idx('u','i,k'),mo('='),sat('a',idx('u','candidate,i,k'))),row(idx('a','i,k'),mo('='),idx('u','i,k'),mo('+'),idx('d','i,k')),mv,mp),
      'sphere':row(mi('d'),par(p),mo('='),norm(row(p,mo('−'),mi('c'))),mo('−'),mi('r')),
      'route':row(mi('J'),mo('='),summand(mi('m'),par(row(norm(row(idx('q','j'),mo('−'),idx('q','j−1'))),mo('+'),idx('λ','z'),abs_m(row(idx('z','j'),mo('−'),idx('z','j−1'))))))),
      'follower':row(idx('u','track'),mo('='),idx('k','p'),par(row(mi('r'),mo('−'),p)),mo('+'),idx('k','d'),par(row(idx('v','ref'),mo('−'),v))),
      'tangent':row(mi('t'),mo('='),mi('g'),mo('−'),par(row(sup(mi('g'),mi('T')),mi('n'))),mi('n')),
      'base':arr(row(idx('u','base'),mo('='),par(row(mn(1),mo('−'),mi('w'))),idx('u','track'),mo('+'),mi('w'),idx('u','obs'),mo('+'),idx('u','pair')),row(mi('w'),mo('='),mn('0.9'),sqrt(mi('ρ')))),
      'repulsion':arr(row(idx('F','n'),mo('='),mn(20),mi('max'),par(seq(mn(0),frac(row(mn(6),mo('−'),mi('d')),mn(6)))),mi('max'),par(seq(mn('0.2'),row(mn(1),mo('−'),sup(mi('e'),row(mo('−'),mn('0.035'),sup(norm(mi('g')),mn(2))))))),mi('n')),
          row(idx('F','t'),mo('='),mn('1.5'),par(row(mn(1),mo('−'),frac(mi('d'),mn(6)))),frac(mi('t'),norm(mi('t'))),mt('   0 < d < 3'))),
      'stereo':row(mi('Z'),mo('='),frac(row(mi('f'),mi('B')),idx('d','disparity'))),
    }
    ep,ev=idx('e','p'),idx('e','v')
    ep2,ev2=sup(norm(ep),mn(2)),sup(norm(ev),mn(2))
    phi,psi=mi('φ'),mi('ψ')
    risks=[sup(idx('ρ',a),mn(2)) for a in ['o','p','b']]
    eq['features']=arr(row(ep,mo('='),frac(row(mi('r'),mo('−'),p),mn(4)),mo(','),ev,mo('='),frac(row(idx('v','ref'),mo('−'),v),mn(3))),
        row(psi,mo('='),vec(frac(ep,sqrt(row(mn(1),mo('+'),ep2))),frac(ev,sqrt(row(mn(1),mo('+'),ev2))))),
        row(phi,mo('='),sup(row(mo('['),seq(ep2,ev2,row(mn(2),sup(ep,mi('T')),ev),*risks),mo(']')),mi('T'))))
    cap=par(seq(mn('0.4'),row(mn('0.25'),norm(sat('a',idx('u','base'))))))
    cap_row=row(mi('c'),mo('='),mi('min'),cap)
    policy_row=row(idx('u','ADP'),mo('='),sat('c',row(mn('0.5'),idx('W','A'),psi)))
    eq['residual']=arr(cap_row,policy_row)
    eq['cost']=row(idx('ℓ','k'),mo('='),dt,par(row(mn('0.7'),ep2,mo('+'),mn('0.35'),ev2,mo('+'),summand(mn(3),sup(idx('ρ','j'),mn(2))),mo('+'),mn('0.03'),sup(norm(idx('u','eff')),mn(2)))))
    wc=idx('w','C');phik=idx('φ','k');phin=sub(phi,k1)
    eq['td']=arr(row(idx('δ','k'),mo('='),idx('ℓ','k'),mo('+'),mn('0.95'),sup(wc,mi('T')),phin,mo('−'),sup(wc,mi('T')),phik),
       row(idx('Δw','C'),mo('='),frac(mn('0.05'),mn(4)),summand(mn(5),frac(row(idx('δ','i'),idx('φ','i,k')),row(mn('0.001'),mo('+'),sup(norm(idx('φ','i,k')),mn(2)))),'i',2)))
    E=idx('E','k');ei=norm(row(idx('p','i,k'),mo('−'),idx('r','i,k')))
    rms=row(idx('E','RMS'),mo('='),sqrt(frac(summand(mi('K'),sup(E,mn(2)),'k'),mi('K'))))
    eq['rms']=arr(row(E,mo('='),sub(mi('max'),row(mi('i'),mo('='),mn(2),mo('…'),mn(5))),ei),rms)
    effort=row(idx('I','u'),mo('='),dt,summand(mi('K'),summand(mn(5),sup(norm(idx('u','i,k')),mn(2)),'i'),'k'))
    slew=row(idx('s','i,k'),mo('='),frac(norm(row(idx('u','i,k'),mo('−'),idx('u','i,k−1'))),dt))
    eq['metrics']=arr(rms,effort,slew)
    return eq

XSL=etree.XSLT(etree.parse(r'C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL'))
EQS=equations()

def set_run(run, size=11.5, bold=False, color='000000', font='宋体'):
    run.font.name='Arial';run.font.size=Pt(size);run.bold=bold;run.font.color.rgb=RGBColor.from_string(color)
    rpr=run._element.get_or_add_rPr();rf=rpr.find(qn('w:rFonts'))
    if rf is None:rf=OxmlElement('w:rFonts');rpr.insert(0,rf)
    for a,b in [('eastAsia',font),('ascii','Arial'),('hAnsi','Arial')]:rf.set(qn('w:'+a),b)

def hyperlink(p,label,path,size=10,color='285477'):
    part=p.part;rid=part.relate_to(Path(path).as_uri(),'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rid)
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr')
    rf=OxmlElement('w:rFonts');rf.set(qn('w:eastAsia'),'宋体');rf.set(qn('w:ascii'),'Arial');pr.append(rf)
    cc=OxmlElement('w:color');cc.set(qn('w:val'),color);pr.append(cc)
    sz=OxmlElement('w:sz');sz.set(qn('w:val'),str(int(size*2)));pr.append(sz)
    r.append(pr);t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)

def math_par(doc,key):
    xml='<math xmlns="http://www.w3.org/1998/Math/MathML" display="block">'+EQS[key]+'</math>'
    result=XSL(etree.fromstring(xml.encode('utf-8')))
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(7)
    p.paragraph_format.keep_together=True
    node=copy.deepcopy(result.getroot())
    for r in node.iter('{http://schemas.openxmlformats.org/officeDocument/2006/math}r'):
        rp=OxmlElement('w:rPr');sz=OxmlElement('w:sz');sz.set(qn('w:val'),'23');rp.append(sz);r.insert(0,rp)
    p._p.append(node)

def table(doc, lines, teacher):
    rows=[[s.strip() for s in ln.strip().strip('|').split('|')] for ln in lines]
    rows=[r for r in rows if not all(re.fullmatch(r':?-+:?',x.replace(' ','')) for x in r)]
    n=len(rows[0]);t=doc.add_table(rows=1, cols=n);t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
    widths=([1.7,5.3] if n==2 else [2,2.5,2.5] if n==3 else [2.5]+[4.5/(n-1)]*(n-1))
    if n>=6:widths=[1.5]+[5.5/(n-1)]*(n-1)
    if rows[0][0]=='阅读顺序':widths=[.7,3.5,2.8]
    if rows[0][0]=='模块':widths=[1.3,2.9,2.8]
    for i,width in enumerate(widths):t.columns[i].width=Inches(width)
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        q=OxmlElement('w:'+edge);q.set(qn('w:val'),'single');q.set(qn('w:sz'),'4');q.set(qn('w:color'),'D9D9D9');borders.append(q)
    pr.append(borders)
    for ri,rowvals in enumerate(rows):
        cells=t.rows[0].cells if ri==0 else t.add_row().cells
        trpr=cells[0]._tc.getparent().get_or_add_trPr()
        if ri==0:
            rep=OxmlElement('w:tblHeader');trpr.append(rep)
        cant=OxmlElement('w:cantSplit');trpr.append(cant)
        for ci,(cell,value) in enumerate(zip(cells,rowvals)):
            cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'354C5B' if ri==0 else ('F2F5F7' if ri%2 else 'FFFFFF'));tcpr.append(shade)
            mar=OxmlElement('w:tcMar')
            for side,val in [('top','75'),('bottom','75'),('left','85'),('right','85')]:
                x=OxmlElement('w:'+side);x.set(qn('w:w'),val);x.set(qn('w:type'),'dxa');mar.append(x)
            tcpr.append(mar)
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1.13
            if ci>0 and (len(value)<17 or re.match(r'^[0-9−±≤≥—]',value)):p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            run=p.add_run(value);set_run(run,10.2 if n<=4 else 9.5,ri==0,'FFFFFF' if ri==0 else '000000','黑体' if ri==0 else '宋体')
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=0.3
    p.add_run().font.size=Pt(2)

def build_doc(src, name, teacher=False):
    text=src.read_text(encoding='utf-8')
    normalized=text.replace('−','-').replace('×10⁻¹³','e-13').replace('10⁻⁶','1e-6').replace('10⁻³','1e-3').replace('10⁻⁸','1e-8').replace('10⁻⁷','1e-7').replace('10⁻⁴','1e-4')
    if normalized!=text:src.write_text(normalized,encoding='utf-8');text=normalized
    doc=Document();s=doc.sections[0]
    s.page_width=Inches(8.5);s.page_height=Inches(11)
    s.top_margin=Inches(.67);s.bottom_margin=Inches(.65)
    s.left_margin=s.right_margin=Inches(.75)
    s.header_distance=Inches(.2);s.footer_distance=Inches(.28)
    body=11.0 if teacher else 11.5
    for st in ['Normal','Body Text']:
        sty=doc.styles[st];sty.font.name='Arial';sty.font.size=Pt(body)
        sty._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')
        sty.paragraph_format.space_after=Pt(7)
        sty.paragraph_format.line_spacing=1.25
        sty.paragraph_format.widow_control=True
    for st,sz in [('Title',23),('Subtitle',13),('Heading 1',16),('Heading 2',12.5)]:
        sty=doc.styles[st];sty.font.name='Arial';sty.font.size=Pt(sz);sty.font.color.rgb=RGBColor(0,0,0)
        sty._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'黑体')
        sty.paragraph_format.space_before=Pt(0 if st in ['Title','Subtitle'] else 4)
        sty.paragraph_format.space_after=Pt(12 if st=='Title' else 10)
    doc.styles['Title'].font.bold=True
    doc.styles['Heading 1'].font.bold=True
    settings=doc.settings.element
    update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
    fp=s.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=fp.add_run();set_run(r,9,False,'666666')
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');r._r.addnext(fld)
    doc.core_properties.title='无人机编队智能避障项目'+('技术报告 教师审阅版' if teacher else '整体报告 学生阅读版')
    doc.core_properties.author='无人机编队智能避障项目'
    doc.core_properties.subject='源码 实验与技术路线整体说明'
    doc.core_properties.comments='2026年10月30日统一Actor重验 证据索引随报告保存'
    chunks=text.split('@page')
    figures_used=[]
    headings=[]
    for pi,chunk in enumerate(chunks):
        if pi:doc.add_page_break()
        lines=chunk.strip().splitlines();i=0
        while i<len(lines):
            ln=lines[i].strip()
            if not ln:i+=1;continue
            if ln.startswith('# '):
                p=doc.add_paragraph(ln[2:],style='Title');headings.append(ln[2:]);i+=1;continue
            if ln.startswith('## '):
                label=ln[3:];headings.append(label)
                doc.add_paragraph(label,style='Heading 1');i+=1;continue
            if ln.startswith('### '):doc.add_paragraph(ln[4:],style='Heading 2');i+=1;continue
            if ln.startswith('@eq '):math_par(doc,ln[4:]);i+=1;continue
            if ln.startswith('@figure '):
                key,caption=map(str.strip,ln[8:].split('|',1));path=FIG/(key+'.png')
                assert path.exists(),path
                p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after=Pt(3);p.paragraph_format.keep_with_next=True
                picture_width=6.8
                if teacher and key=='噪声控制与轨迹':picture_width=6.0
                elif key=='安全层同图轨迹':picture_width=6.35
                elif teacher and key=='控制信息流':picture_width=6.35
                p.add_run().add_picture(str(path),width=Inches(picture_width))
                p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(9)
                set_run(p.add_run(caption),9.0,False,'555555')
                figures_used.append(key);i+=1;continue
            if ln.startswith('@source '):
                p=doc.paragraphs[-1] if doc.paragraphs and doc.paragraphs[-1].text else doc.add_paragraph()
                set_run(p.add_run('  依据  '),9,False,'666666')
                for key in ln[8:].split():
                    label,path=EVIDENCE[key];hyperlink(p,key,path,9);set_run(p.add_run('   '),9)
                i+=1;continue
            if ln=='@evidence':
                # Compact bibliography; source labels are active hyperlinks.
                for key,(label,path) in EVIDENCE.items():
                    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(5);p.paragraph_format.line_spacing=1.12
                    set_run(p.add_run(key+'  '),9.4,True)
                    hyperlink(p,label,path,9.4)
                i+=1;continue
            if ln.startswith('|'):
                block=[]
                while i<len(lines) and lines[i].strip().startswith('|'):block.append(lines[i].strip());i+=1
                table(doc,block,teacher);continue
            # Every authored paragraph remains one Word paragraph.
            p=doc.add_paragraph()
            isnumbered=bool(re.match(r'^\d+\. ',ln))
            if isnumbered:
                p.paragraph_format.left_indent=Inches(.14);p.paragraph_format.first_line_indent=Inches(-.14)
            set_run(p.add_run(ln),body)
            i+=1
    dst=ROOT/(name+'.docx');doc.save(dst)
    summary={'file':str(dst),'authored_pages':len(chunks),'headings':headings,'figures':figures_used,'characters':len(text),'native_equations':len(re.findall(r'^@eq ',text,re.M))}
    (ROOT/'_work'/(name+'_structure.json')).write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(summary|{'headings':len(headings)},ensure_ascii=False))
    return dst

def main():
    (ROOT/'_work').mkdir(exist_ok=True)
    evidence=[]
    for key,(label,path) in EVIDENCE.items():
        assert path.exists(),path
        entry={'id':key,'label':label,'path':str(path)}
        if path.is_file():entry['sha256']=hashlib.sha256(path.read_bytes()).hexdigest()
        evidence.append(entry)
    for path in [ROBUST/'逐任务汇总.csv',ROBUST/'恢复记录.csv',DYNAMIC/'逐任务汇总.csv',SAFETY/'任务结局汇总.csv']:
        evidence.append({'label':'图表与统计输入','path':str(path),'sha256':hashlib.sha256(path.read_bytes()).hexdigest()})
    (ROOT/'evidence_index.json').write_text(json.dumps(evidence,ensure_ascii=False,indent=2),encoding='utf-8')
    facts=make_figures()
    build_doc(ROOT/'学生版_正文.md','无人机编队项目整体报告_学生版')
    build_doc(ROOT/'教师版_正文.md','无人机编队项目技术报告_教师版',True)
    print(json.dumps(facts,ensure_ascii=False))

if __name__=='__main__':main()
