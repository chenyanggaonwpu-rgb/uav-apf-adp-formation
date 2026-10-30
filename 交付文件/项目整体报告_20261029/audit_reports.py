from pathlib import Path
import json,re,hashlib,unicodedata,collections
from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader
ROOT=Path(__file__).resolve().parent;P=next(p for p in ROOT.parents if (p/'unified_actor3d').is_dir());O=P/'results/unified_actor_20261029'
def norm(s):return re.sub(r'\s+','',unicodedata.normalize('NFKC',s))
def main():
    complete=json.loads((O/'completion.json').read_text(encoding='utf-8'));assert complete['trials']==3840 and complete['allAuditsPassed']
    receipts=json.loads((ROOT/'_work/offline_export_receipt.json').read_text(encoding='utf-8'));raster=json.loads((ROOT/'_work/render_receipt.json').read_text(encoding='utf-8'))
    results=[]
    for index,name in enumerate(['无人机编队项目整体报告_学生版','无人机编队项目技术报告_教师版']):
        d=Document(ROOT/(name+'.docx'));reader=PdfReader(ROOT/(name+'.pdf'));pdf=''
        for i,page in enumerate(reader.pages):
            text=page.extract_text() or '';text=re.sub(r'^'+str(i+1)+r'\s*','',text)
            assert len(text.strip())>50,(name,'orphan page',i+1,text)
            assert float(page.mediabox.width)==612 and float(page.mediabox.height)==792
            pdf+=text
        pdf=norm(pdf);miss=[]
        for paragraph in d.element.body.iter(qn('w:p')):
            if list(paragraph.iter(qn('m:oMath'))):continue
            text=''.join(t.text or '' for t in paragraph.iter(qn('w:t')))
            if text.strip() and norm(text) not in pdf:miss.append(text)
        assert not miss,(name,'Word paragraphs missing from PDF',miss)
        xml=d.element.xml;equations=len(d.element.xpath('.//m:oMathPara'))
        assert equations==[8,11][index]
        assert len(d.inline_shapes)==6
        assert '2.3.4-unified-actor-20261029' in pdf and '6467' in pdf
        assert '3840' in pdf and norm('基础方案成功 71 次') in pdf and norm('Actor 方案成功 122 次') in pdf and norm('规则脱困成功 124 次') in pdf
        assert all(t._tbl.tblPr.find(qn('w:tblBorders')) is not None for t in d.tables)
        assert all(not p._p.xpath('./w:pPr/w:pBdr') for p in d.paragraphs if p.style.name in ['Title','Heading 1','Heading 2'])
        captions=[p.text for p in d.paragraphs if re.match(r'^图\d+',p.text)]
        assert [int(re.match(r'^图(\d+)',s).group(1)) for s in captions]==list(range(1,7)),captions
        assert receipts[index]['native_omml_equivalence'] and receipts[index]['docx_sha256']==hashlib.sha256((ROOT/(name+'.docx')).read_bytes()).hexdigest()
        assert len(raster[index]['pages'])==len(reader.pages)
        assert len(list((ROOT/'_work'/('render_student' if index==0 else 'render_teacher')).glob('page-*.png')))==len(reader.pages)
        results.append(dict(file=name,pages=len(reader.pages),equations=equations,figures=6,tables=len(d.tables),allWordTextPresentInPDF=True,nativeOMMLEquivalent=True,
            docxSha256=hashlib.sha256((ROOT/(name+'.docx')).read_bytes()).hexdigest(),pdfSha256=hashlib.sha256((ROOT/(name+'.pdf')).read_bytes()).hexdigest()))
    rows=[json.loads(p.read_text(encoding='utf-8')) for p in (O/'records').glob('*.json')]
    assert len(rows)==3840 and len({r['id'] for r in rows})==3840
    assert len({r['modelSha256'] for r in rows})==1 and len({r['runtimeSourceSha256'] for r in rows})==1
    assert all(r['legacyADPCalls']==0 and r['actorLoaded'] and not r['trainingPerformed'] for r in rows)
    assert collections.Counter(r['campaign'] for r in rows)==dict(repair=616,dynamic=200,robust=1824,integration=256,factorial=864,single=80)
    geometry=json.loads((O/'dynamic_geometry_audit.json').read_text(encoding='utf-8'));assert len(geometry)==200 and all(r['passed'] for r in geometry)
    result=dict(passed=True,reports=results,trials=3840,physicsSteps=complete['steps'],independentAudits=3840,dynamicGeometryAudits=200,
        originalExecutionErrorsPreserved=complete['executionErrorAttempts'],visualInspection='See visual_qa.json for the complete page-by-page review')
    (ROOT/'交付复核.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps(result,ensure_ascii=False))
if __name__=='__main__':main()
